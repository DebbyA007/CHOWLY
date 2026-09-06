import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import content as C
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT, ERD = sys.argv[1], sys.argv[2]
INK = RGBColor(0x1a, 0x1a, 0x1a)
MUTED = RGBColor(0x5a, 0x5a, 0x5a)

doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin = Cm(2.0); sec.bottom_margin = Cm(1.8)

# Base styles: one family, sized like the PDF so the two read the same.
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.font.color.rgb = INK
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.15
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
for name, size, before, after in (("Heading 1", 17, 20, 10), ("Heading 2", 13.5, 16, 7), ("Heading 3", 11.5, 12, 5)):
    s = doc.styles[name]
    s.font.name = "Calibri"; s.font.size = Pt(size); s.font.bold = True
    s.font.color.rgb = INK
    s.paragraph_format.space_before = Pt(before); s.paragraph_format.space_after = Pt(after)

def field(paragraph, instr, cached_runs=None):
    """A real Word field. Word recomputes it; a viewer that will not shows the cached text."""
    r = paragraph.add_run()._r
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "begin"); fc.set(qn("w:dirty"), "true")
    r.append(fc)
    r2 = paragraph.add_run()._r
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    r2.append(it)
    r3 = paragraph.add_run()._r
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    r3.append(sep)
    if cached_runs:
        cached_runs(paragraph)
    r4 = paragraph.add_run()._r
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    r4.append(end)

# ---- title page
def line(text, size=10.5, bold=False, colour=INK, after=4, align=None):
    p = doc.add_paragraph()
    if align: p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    run = p.add_run(text); run.bold = bold
    run.font.size = Pt(size); run.font.color.rgb = colour
    return p

doc.add_paragraph().paragraph_format.space_after = Pt(120)
line(C.TITLE, size=30, bold=True, after=4)
line(C.SUBTITLE, size=13, colour=MUTED, after=22)
line(C.AUTHOR, size=11, bold=True, after=3)
line(f"Repository   {C.REPO}", size=10, colour=MUTED, after=2)
line(f"Deployed     {C.LIVE}", size=10, colour=MUTED, after=26)
line("This document answers the three deliverables of the CHOWLY build assignment in the order the brief lists them. It is written to be read without any other document open.",
     size=10.5, colour=MUTED)
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ---- contents: a real TOC field, with the section list cached inside it so a viewer
# that does not recompute fields still shows something useful.
h = doc.add_paragraph("Contents"); h.style = doc.styles["Heading 1"]
toc_p = doc.add_paragraph()
def cached(paragraph):
    for kind, val in C.DOC:
        if kind == "h1":
            r = paragraph.add_run("\n" + val); r.bold = True; r.font.size = Pt(11)
        elif kind == "h2":
            r = paragraph.add_run("\n    " + val); r.font.size = Pt(10)
        elif kind == "h3":
            r = paragraph.add_run("\n        " + val); r.font.size = Pt(9.5); r.font.color.rgb = MUTED
    paragraph.add_run("\n")
field(toc_p, r'TOC \o "1-3" \h \z \u ', cached)
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ---- footer: page number field
footer = sec.footer
fp = footer.paragraphs[0]
fp.text = "CHOWLY: submission document"
fp.paragraph_format.tab_stops.add_tab_stop(sec.page_width - sec.left_margin - sec.right_margin)
fp.add_run("\t")
field(fp, " PAGE ", lambda p: p.add_run("1"))
for r in fp.runs:
    r.font.size = Pt(8); r.font.color.rgb = MUTED

# ---- body
num = 0
for kind, val in C.DOC:
    if kind in ("h1", "h2", "h3"):
        doc.add_paragraph(val, style=doc.styles[{"h1": "Heading 1", "h2": "Heading 2", "h3": "Heading 3"}[kind]])
        num = 0
    elif kind == "p":
        p = doc.add_paragraph(val); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; num = 0
    elif kind == "bullet":
        p = doc.add_paragraph(val, style="List Bullet"); p.paragraph_format.space_after = Pt(5)
    elif kind == "num":
        num += 1
        p = doc.add_paragraph(val, style="List Number"); p.paragraph_format.space_after = Pt(5)
    elif kind == "image":
        doc.add_picture(ERD, width=Cm(16.6))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = doc.add_paragraph()
        r = cap.add_run("The data model as implemented. Every entity of the engineered model is present; the fields added or changed by the fifteen deltas are visible on Menu, MenuItem and Order.")
        r.font.size = Pt(8.5); r.font.color.rgb = MUTED
    elif kind == "table":
        head, rows = val
        t = doc.add_table(rows=1, cols=len(head))
        t.style = "Table Grid"
        for i, htxt in enumerate(head):
            c = t.rows[0].cells[i]; c.text = ""
            run = c.paragraphs[0].add_run(htxt); run.bold = True; run.font.size = Pt(9)
            shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "F2F0EC")
            c._tc.get_or_add_tcPr().append(shade)
        for row in rows:
            cells = t.add_row().cells
            for i, txt in enumerate(row):
                cells[i].text = ""
                run = cells[i].paragraphs[0].add_run(str(txt)); run.font.size = Pt(9)
        # repeat the header row across a page break
        trPr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

# the Word file's own author property, not only the name on the page
doc.core_properties.author = C.AUTHOR
doc.core_properties.title = "CHOWLY: submission document"
doc.core_properties.comments = ""
doc.save(OUT)
print("wrote", OUT)
